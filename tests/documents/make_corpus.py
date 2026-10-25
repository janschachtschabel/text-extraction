"""Builds the document corpus in ``tests/documents/corpus/``. Dev only; the generated files are
committed, so the tests do not need this script, reportlab, python-docx or Pillow at run time.

    uv run python tests/documents/make_corpus.py

The output is deterministic (fixed dates and ids), so a re-run changes a file only when this
script or the system font changes. Later tasks add their files here.
"""

import datetime
import io
import sys
import zipfile
from pathlib import Path

import docx
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import pdfencrypt
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

CORPUS = Path(__file__).parent / "corpus"

_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
)

GERMAN_PAGES = (
    (
        "Die Photosynthese",
        "Pflanzen wandeln Lichtenergie in chemische Energie um. Dabei nehmen die Blätter",
        "Kohlenstoffdioxid aus der Luft auf und geben Sauerstoff ab. Der Vorgang läuft",
        "in den Chloroplasten ab und braucht Wasser, das die Wurzeln aus dem Boden holen.",
        "Größere Pflanzen speichern den entstandenen Zucker als Stärke in Knollen.",
    ),
    (
        "Wichtige Stoffe",
        "Ohne Chlorophyll, den grünen Farbstoff der Blätter, könnte keine Pflanze das",
        "Sonnenlicht nutzen. Über die Spaltöffnungen tauschen die Blätter Gase mit der",
        "Umgebung aus. Im Herbst bauen Bäume das Chlorophyll ab; dann färben sich die",
        "Blätter gelb und rot. Die Übung zeigt: Ohne Licht wächst keine grüne Pflanze.",
    ),
)

TABLE_ROWS = (
    ("Stoff", "Menge", "Aufgabe"),
    ("Wasser", "6 Teile", "Rohstoff"),
    ("Kohlenstoffdioxid", "6 Teile", "Rohstoff"),
    ("Traubenzucker", "1 Teil", "Produkt"),
)


def _text_pdf() -> bytes:
    """Three pages of German text with a ruled table; title, author and keywords in the info
    dictionary (reportlab's ``invariant`` mode fixes its dates to 2000-01-01)."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1, pageCompression=0)
    pdf.setTitle("Photosynthese im Überblick")
    pdf.setAuthor("Erika Mustermann")
    pdf.setSubject("Biologie, Klasse 7")
    pdf.setKeywords("Photosynthese, Chlorophyll; Pflanzen")
    for lines in GERMAN_PAGES:
        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(72, 770, lines[0])
        pdf.setFont("Helvetica", 11)
        for number, line in enumerate(lines[1:]):
            pdf.drawString(72, 740 - 18 * number, line)
        pdf.showPage()
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(72, 770, "Die Gleichung")
    pdf.setFont("Helvetica", 11)
    top, row_height, widths = 720, 24, (150, 90, 120)
    for row_number, row in enumerate(TABLE_ROWS):
        y = top - row_number * row_height
        x = 72
        for cell, width in zip(row, widths, strict=True):
            pdf.rect(x, y - row_height, width, row_height)
            pdf.drawString(x + 6, y - 16, cell)
            x += width
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _font_path() -> str:
    for candidate in _FONT_CANDIDATES:
        if Path(candidate).is_file():
            return candidate
    sys.exit("DejaVuSans.ttf not found: install fonts-dejavu-core or add its path")


def _scan_pages() -> list[Image.Image]:
    font = ImageFont.truetype(_font_path(), 30)
    title_font = ImageFont.truetype(_font_path(), 40)
    pages = []
    for lines in GERMAN_PAGES:
        image = Image.new("L", (1240, 600), 255)
        draw = ImageDraw.Draw(image)
        draw.text((60, 50), lines[0], font=title_font, fill=0)
        for number, line in enumerate(lines[1:]):
            draw.text((60, 140 + 60 * number), line, font=font, fill=0)
        pages.append(image)
    return pages


def _scan_pdf(pages: list[Image.Image]) -> bytes:
    """Two pages that are one picture each and carry no text layer."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1)
    pdf.setTitle("Scan")
    for image in pages:
        width, height = A4
        pdf.drawImage(ImageReader(image), 0, height - width * 600 / 1240, width, width * 600 / 1240)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _encrypted_pdf() -> bytes:
    buffer = io.BytesIO()
    encryption = pdfencrypt.StandardEncryption(
        userPassword="geheim", ownerPassword="eigentümer", strength=128
    )
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1, encrypt=encryption)
    pdf.setFont("Helvetica", 12)
    pdf.drawString(72, 770, "Dieses Dokument verlangt ein Kennwort.")
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _raw_pdf(objects: list[bytes], root: int) -> bytes:
    """A PDF from object bodies (object numbers 1..n in order), with a plain xref table: the
    shapes that reportlab cannot write (no pages at all, an XMP stream)."""
    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        root,
        xref,
    )
    return bytes(out)


def _stream(dictionary: bytes, data: bytes) -> bytes:
    return b"<< " + dictionary + b" /Length %d >>\nstream\n" % len(data) + data + b"\nendstream"


def _empty_pdf() -> bytes:
    return _raw_pdf(
        [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [] /Count 0 >>"], root=1
    )


_XMP = """<?xpacket begin="﻿" id="W5M0MpCehiHzreSzNTczkc9d"?>
<x:xmpmeta xmlns:x="adobe:ns:meta/">
 <rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">
  <rdf:Description rdf:about=""
    xmlns:dc="http://purl.org/dc/elements/1.1/"
    xmlns:cc="http://creativecommons.org/ns#"
    xmlns:xmp="http://ns.adobe.com/xap/1.0/"
    xmlns:xmpRights="http://ns.adobe.com/xap/1.0/rights/">
   <dc:title>
    <rdf:Alt>
     <rdf:li xml:lang="x-default">Lizenzierte Handreichung</rdf:li>
    </rdf:Alt>
   </dc:title>
   <dc:creator><rdf:Seq><rdf:li>Erika Mustermann</rdf:li></rdf:Seq></dc:creator>
   <dc:description>
    <rdf:Alt>
     <rdf:li xml:lang="x-default">Ein frei lizenziertes Arbeitsblatt.</rdf:li>
    </rdf:Alt>
   </dc:description>
   <dc:rights>
    <rdf:Alt>
     <rdf:li xml:lang="x-default">Copyright 2024 Erika Mustermann</rdf:li>
    </rdf:Alt>
   </dc:rights>
   <dc:language><rdf:Bag><rdf:li>de-DE</rdf:li></rdf:Bag></dc:language>
   <dc:subject><rdf:Bag><rdf:li>Biologie</rdf:li><rdf:li>OER</rdf:li></rdf:Bag></dc:subject>
   <cc:license rdf:resource="https://creativecommons.org/licenses/by-sa/4.0/"/>
   <xmpRights:WebStatement>https://creativecommons.org/licenses/by-sa/4.0/</xmpRights:WebStatement>
   <xmp:CreateDate>2024-03-15T09:30:00+01:00</xmp:CreateDate>
   <xmp:ModifyDate>2024-04-01T12:00:00Z</xmp:ModifyDate>
  </rdf:Description>
 </rdf:RDF>
</x:xmpmeta>
<?xpacket end="w"?>"""


def _xmp_pdf() -> bytes:
    content = (
        b"BT /F1 12 Tf 72 770 Td (Dieses Arbeitsblatt steht unter der Lizenz CC BY-SA 4.0. "
        b"Es darf mit Namensnennung weitergegeben werden.) Tj ET"
    )
    return _raw_pdf(
        [
            b"<< /Type /Catalog /Pages 2 0 R /Metadata 5 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R"
            b" /Resources << /Font << /F1 6 0 R >> >> >>",
            _stream(b"", content),
            _stream(b"/Type /Metadata /Subtype /XML", _XMP.encode()),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        ],
        root=1,
    )


def _fix_zip_dates(blob: bytes) -> bytes:
    """The same archive with every entry dated 2000-01-01, so that a re-run gives the same bytes."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(blob)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=(2000, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = info.external_attr
            target.writestr(entry, source.read(info.filename))
    return out.getvalue()


def _core_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Photosynthese", level=1)
    document.add_paragraph("Pflanzen wandeln Lichtenergie in chemische Energie um.")
    properties = document.core_properties
    properties.title = "Photosynthese im Überblick"
    properties.author = "Erika Mustermann"
    properties.subject = "Biologie"
    properties.keywords = "Photosynthese, Chlorophyll; Pflanzen"
    properties.comments = "Ein Arbeitsblatt für die siebte Klasse."
    properties.language = "de-DE"
    properties.created = datetime.datetime(2024, 3, 15, 9, 30, tzinfo=datetime.UTC)
    properties.modified = datetime.datetime(2024, 4, 1, 12, 0, tzinfo=datetime.UTC)
    properties.last_modified_by = "Max Beispiel"
    buffer = io.BytesIO()
    document.save(buffer)
    return _fix_zip_dates(buffer.getvalue())


def main() -> None:
    CORPUS.mkdir(parents=True, exist_ok=True)
    pages = _scan_pages()
    files = {
        "text.pdf": _text_pdf(),
        "scan.pdf": _scan_pdf(pages),
        "encrypted.pdf": _encrypted_pdf(),
        "empty.pdf": _empty_pdf(),
        "xmp-cc-by-sa.pdf": _xmp_pdf(),
        "core.docx": _core_docx(),
        "scan.expected.txt": (
            "\n\n".join("\n".join(lines) for lines in GERMAN_PAGES) + "\n"
        ).encode(),
    }
    for name, data in files.items():
        (CORPUS / name).write_bytes(data)
        print(f"{name}: {len(data)} bytes")


if __name__ == "__main__":
    main()
